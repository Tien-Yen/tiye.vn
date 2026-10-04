from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as _Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xlsm"}
TEMP_FILE_PREFIXES = ("~$",)
MAJOR_SECTION_RE = re.compile(r"^[IVXLC]+[\.)]\s+", re.IGNORECASE)


def clean_text(value) -> str:
    """Normalize Unicode and whitespace without removing Vietnamese characters."""
    if value is None:
        return ""

    text = unicodedata.normalize("NFC", str(value))
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    cleaned_lines: list[str] = []
    for line in text.split("\n"):
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line:
            cleaned_lines.append(line)

    return "\n".join(cleaned_lines)


def file_sha1(path: Path) -> str:
    digest = hashlib.sha1()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def make_block_id(file_id: str, block_index: int) -> str:
    return f"{file_id}-b{block_index:05d}"


def make_file_id(relative_path: str) -> str:
    """Stable ID per source path; avoids collisions between identical copied files."""
    raw = relative_path.replace("\\", "/").lower().encode("utf-8")
    return hashlib.sha1(raw).hexdigest()[:16]


def iter_block_items(parent) -> Iterator[Paragraph | Table]:
    """Yield DOCX paragraphs and tables in their actual document order."""
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise TypeError(f"Unsupported parent type: {type(parent)!r}")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def read_pdf(path: Path) -> list[dict]:
    blocks: list[dict] = []
    reader = PdfReader(str(path))

    for page_number, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        if not text:
            continue

        blocks.append(
            {
                "block_type": "page",
                "text": text,
                "location": {"page": page_number},
            }
        )

    return blocks


def read_docx(path: Path) -> list[dict]:
    """Extract paragraphs and tables without assuming a table has headers."""
    document = Document(str(path))
    blocks: list[dict] = []
    current_section = "Phần mở đầu"

    for item_index, item in enumerate(iter_block_items(document), start=1):
        if isinstance(item, Paragraph):
            text = clean_text(item.text)
            if not text:
                continue

            if MAJOR_SECTION_RE.match(text):
                current_section = text

            blocks.append(
                {
                    "block_type": "paragraph",
                    "text": text,
                    "location": {
                        "section": current_section,
                        "item_index": item_index,
                        "style": clean_text(item.style.name if item.style else ""),
                    },
                }
            )
            continue

        # Table: preserve every row, including one-row / one-column tables.
        rows: list[str] = []
        for row_number, row in enumerate(item.rows, start=1):
            values = [clean_text(cell.text) for cell in row.cells]
            if not any(values):
                continue
            rows.append(
                f"Dòng {row_number}: "
                + " | ".join(
                    f"Ô {col_number}: {value}"
                    for col_number, value in enumerate(values, start=1)
                    if value
                )
            )

        if not rows:
            continue

        blocks.append(
            {
                "block_type": "table",
                "text": "\n".join(rows),
                "location": {
                    "section": current_section,
                    "item_index": item_index,
                    "table_number": sum(1 for b in blocks if b["block_type"] == "table") + 1,
                },
            }
        )

    return blocks


def read_xlsx(path: Path) -> list[dict]:
    """Extract every non-empty worksheet row with sheet/row/cell metadata."""
    blocks: list[dict] = []
    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )

    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                values: list[tuple[str, str]] = []
                for cell in row:
                    value = clean_text(cell.value)
                    if value:
                        values.append((cell.coordinate, value))

                if not values:
                    continue

                text = "\n".join(
                    [
                        f"Trang tính: {sheet.title}",
                        f"Dòng: {row_number}",
                        " | ".join(
                            f"{coord}={value}" for coord, value in values
                        ),
                    ]
                )

                blocks.append(
                    {
                        "block_type": "row",
                        "text": text,
                        "location": {
                            "sheet": sheet.title,
                            "row": row_number,
                        },
                    }
                )
    finally:
        workbook.close()

    return blocks


def read_file(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix == ".docx":
        return read_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise ValueError(f"Unsupported file type: {suffix}")


def discover_files(source_dir: Path) -> list[Path]:
    files: list[Path] = []
    for path in source_dir.rglob("*"):
        if not path.is_file():
            continue
        if path.name.startswith(TEMP_FILE_PREFIXES):
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        files.append(path)
    return sorted(files, key=lambda p: str(p).lower())


def build_knowledge(source_dir: Path) -> tuple[list[dict], list[dict]]:
    all_blocks: list[dict] = []
    manifest: list[dict] = []

    files = discover_files(source_dir)
    print(f"Tìm thấy {len(files)} tài liệu được hỗ trợ.")

    seen_hashes: dict[str, list[str]] = {}

    for file_path in files:
        relative_path = file_path.relative_to(source_dir).as_posix()
        content_sha1 = file_sha1(file_path)
        file_id = make_file_id(relative_path)
        suffix = file_path.suffix.lower()

        print(f"\nĐọc: {relative_path}")

        seen_hashes.setdefault(content_sha1, []).append(relative_path)

        entry = {
            "file_id": file_id,
            "source": file_path.name,
            "relative_path": relative_path,
            "file_type": suffix.lstrip("."),
            "content_sha1": content_sha1,
            "content_sha1": content_sha1,
            "blocks": 0,
            "errors": [],
        }

        try:
            blocks = read_file(file_path)
            for index, block in enumerate(blocks):
                record = {
                    "id": make_block_id(file_id, index),
                    "file_id": file_id,
                    "source": file_path.name,
                    "relative_path": relative_path,
                    "file_type": suffix.lstrip("."),
                    "block_index": index,
                    "block_type": block["block_type"],
                    "text": block["text"],
                    "location": block["location"],
                }
                all_blocks.append(record)

            entry["blocks"] = len(blocks)
            print(f"  → {len(blocks)} blocks")

        except Exception as exc:
            entry["errors"].append(f"{type(exc).__name__}: {exc}")
            print(f"  [LỖI] {type(exc).__name__}: {exc}")

        manifest.append(entry)

    # Chỉ cảnh báo bản sao giống hệt nhau; không tự xóa hoặc gộp tài liệu.
    duplicate_groups = [paths for paths in seen_hashes.values() if len(paths) > 1]
    if duplicate_groups:
        print("\n[CẢNH BÁO] Phát hiện file có nội dung giống hệt nhau:")
        for group in duplicate_groups:
            for item in group:
                print(f"  - {item}")

    return all_blocks, manifest


def save_jsonl(records: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")


def save_json(data, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(data, handle, ensure_ascii=False, indent=2)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="School AI - Giai đoạn 1: đọc và chuẩn hóa tài liệu, chưa tạo embedding."
    )
    parser.add_argument(
        "--source",
        default="data/documents",
        help="Thư mục tài liệu nguồn. Mặc định: data/documents",
    )
    parser.add_argument(
        "--output",
        default="data/processed",
        help="Thư mục dữ liệu đã chuẩn hóa. Mặc định: data/processed",
    )
    args = parser.parse_args()

    source_dir = Path(args.source)
    output_dir = Path(args.output)

    if not source_dir.exists():
        raise SystemExit(f"Không tìm thấy thư mục tài liệu: {source_dir}")

    print("=" * 60)
    print("SCHOOL AI - BƯỚC 1: DOCUMENT INGESTION")
    print("Chưa tạo embedding / Chroma / gọi AI")
    print("=" * 60)

    blocks, manifest = build_knowledge(source_dir)

    if not blocks:
        print("\nKhông đọc được block nào từ tài liệu.")
        raise SystemExit(1)

    blocks_path = output_dir / "blocks.jsonl"
    manifest_path = output_dir / "manifest.json"
    save_jsonl(blocks, blocks_path)
    save_json(manifest, manifest_path)

    counts: dict[str, int] = {}
    block_type_counts: dict[str, int] = {}
    error_count = 0

    for entry in manifest:
        counts[entry["file_type"]] = counts.get(entry["file_type"], 0) + entry["blocks"]
        error_count += len(entry["errors"])

    for block in blocks:
        block_type = block["block_type"]
        block_type_counts[block_type] = block_type_counts.get(block_type, 0) + 1

    print("\n" + "=" * 60)
    print("KẾT QUẢ INGESTION")
    print("=" * 60)
    print(f"Tổng file: {len(manifest)}")
    print(f"Tổng block: {len(blocks)}")
    for file_type, count in sorted(counts.items()):
        print(f"  {file_type}: {count} blocks")
    print("Theo loại block:")
    for block_type, count in sorted(block_type_counts.items()):
        print(f"  {block_type}: {count}")
    print(f"File có lỗi: {error_count}")
    print(f"Đã lưu: {blocks_path}")
    print(f"Đã lưu: {manifest_path}")

    print("\nKiểm tra nhanh nội dung:")
    for phrase in ("Trương Thị Thanh Châm", "HIỆU TRƯỞNG", "Tien Yen'got talent"):
        matches = [b for b in blocks if phrase.lower() in b["text"].lower()]
        print(f"  {phrase!r}: {len(matches)} block")


if __name__ == "__main__":
    main()
