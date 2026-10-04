from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Text chunks for prose documents.
MAX_CHARS = 1400
MAX_TABLE_CHARS = 5000

SECTION_RE = re.compile(r"^\s*[IVXLC]+[\.)]\s+", re.IGNORECASE)


def load_jsonl(path: Path) -> list[dict]:
    records: list[dict] = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                records.append(json.loads(line))
    return records


def clean(text: Any) -> str:
    if text is None:
        return ""

    lines: list[str] = []
    for line in str(text).replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def pack_lines(lines: list[str], max_chars: int) -> list[str]:
    chunks: list[str] = []
    current: list[str] = []
    current_len = 0

    for raw_line in lines:
        line = clean(raw_line)
        if not line:
            continue

        add_len = len(line) + (1 if current else 0)
        if current and current_len + add_len > max_chars:
            chunks.append("\n".join(current))
            current = [line]
            current_len = len(line)
        else:
            current.append(line)
            current_len += add_len

    if current:
        chunks.append("\n".join(current))

    return chunks


def block_to_parts(block: dict) -> list[str]:
    text = clean(block["text"])
    if not text:
        return []

    block_type = block["block_type"]

    if block_type == "table":
        if len(text) <= MAX_TABLE_CHARS:
            return [text]
        return pack_lines(text.split("\n"), MAX_TABLE_CHARS)

    if len(text) <= MAX_CHARS:
        return [text]

    return pack_lines(text.split("\n"), MAX_CHARS)


def make_chunk(
    blocks: list[dict],
    chunk_index: int,
    text: str,
    extra_location: dict | None = None,
) -> dict:
    first = blocks[0]
    location = dict(first.get("location", {}))
    if extra_location:
        location.update(extra_location)

    cleaned_text = clean(text)

    # Nội dung dùng cho retrieval được làm giàu bằng tên tài liệu và section.
    # `text` vẫn giữ nguyên để Qwen đọc đúng dữ liệu gốc, không bị trộn metadata.
    source = str(first.get("source", ""))
    section = str(location.get("section", ""))
    retrieval_parts = [f"Tên tài liệu: {source}"]
    if section:
        retrieval_parts.append(f"Phần: {section}")
    retrieval_parts.append(f"Nội dung:\n{cleaned_text}")
    retrieval_text = "\n".join(retrieval_parts)

    return {
        "id": f"{first['file_id']}-c{chunk_index:05d}",
        "file_id": first["file_id"],
        "source": first["source"],
        "relative_path": first["relative_path"],
        "file_type": first["file_type"],
        "chunk_index": chunk_index,
        "block_ids": [b["id"] for b in blocks],
        "block_types": sorted({b["block_type"] for b in blocks}),
        "location": location,
        "text": cleaned_text,
        "retrieval_text": retrieval_text,
        "char_count": len(cleaned_text),
    }


def chunk_docx(blocks: list[dict]) -> list[dict]:
    chunks: list[dict] = []
    buffer_blocks: list[dict] = []
    buffer_parts: list[str] = []
    buffer_section: str | None = None

    def flush() -> None:
        nonlocal buffer_blocks, buffer_parts, buffer_section
        if not buffer_parts:
            return

        chunks.append(
            make_chunk(
                buffer_blocks,
                len(chunks),
                "\n".join(buffer_parts),
            )
        )
        buffer_blocks = []
        buffer_parts = []
        buffer_section = None

    for block in blocks:
        section = block.get("location", {}).get("section", "")
        block_type = block["block_type"]
        text = clean(block["text"])
        if not text:
            continue

        # Keep tables intact. They often contain the exact answer to a question.
        if block_type == "table":
            flush()
            for part in block_to_parts(block):
                chunks.append(make_chunk([block], len(chunks), part))
            continue

        # A major section heading starts a new semantic area, but the heading
        # itself is NOT useful evidence on its own. Keep it together with the
        # following content so retrieval can match the actual facts under the
        # section instead of returning a heading-only chunk.
        if SECTION_RE.match(text):
            flush()
            buffer_section = section
            buffer_blocks = [block]
            buffer_parts = [text]
            continue

        for part in block_to_parts(block):
            if buffer_parts and buffer_section != section:
                flush()

            proposed = "\n".join(buffer_parts + [part]) if buffer_parts else part
            if buffer_parts and len(proposed) > MAX_CHARS:
                flush()

            if buffer_section is None:
                buffer_section = section

            buffer_blocks.append(block)
            buffer_parts.append(part)

    flush()
    return chunks


def chunk_pdf(blocks: list[dict]) -> list[dict]:
    chunks: list[dict] = []
    for block in blocks:
        for part in block_to_parts(block):
            chunks.append(make_chunk([block], len(chunks), part))
    return chunks


def _xlsx_source_path(blocks: list[dict]) -> Path | None:
    if not blocks:
        return None

    relative_path = blocks[0].get("relative_path", "")
    if not relative_path:
        return None

    # The project layout is fixed by the ingestion stage.
    candidate = Path("data/documents") / relative_path
    if candidate.exists():
        return candidate

    return None


def _excel_text(value: Any) -> str:
    return clean(value).replace("\n", " ")



def _is_class_name(value: str) -> bool:
    return bool(re.fullmatch(r"(?:6|7|8|9|10|11|12)A\d+", value.strip(), re.IGNORECASE))


def _day_display(day: str) -> str:
    names = {
        "2": "Thứ Hai (2)",
        "3": "Thứ Ba (3)",
        "4": "Thứ Tư (4)",
        "5": "Thứ Năm (5)",
        "6": "Thứ Sáu (6)",
        "7": "Thứ Bảy (7)",
    }
    return names.get(str(day), str(day))


def _session_display(session: str) -> str:
    return {
        "S": "Sáng (S)",
        "C": "Chiều (C)",
    }.get(str(session).upper(), str(session))


def _number_key(value: str) -> tuple[int, str]:
    match = re.search(r"\d+", str(value))
    return (int(match.group(0)) if match else 999, str(value))


def _parse_day_value(value: str) -> str:
    match = re.search(r"(?:thứ|thu)\s*(?:hai|ba|tư|bon|năm|nam|sáu|sau|bảy|bay|2|3|4|5|6|7)", value.lower())
    if match:
        text = match.group(0).lower()
        map_name = {
            "hai": "2", "ba": "3", "tư": "4", "bon": "4",
            "năm": "5", "nam": "5", "sáu": "6", "sau": "6",
            "bảy": "7", "bay": "7",
        }
        for name, number in map_name.items():
            if name in text:
                return number
    match = re.search(r"\b([2-7])\b", value)
    return match.group(1) if match else value.strip()


def _parse_session_value(value: str) -> str:
    normalized = value.strip().upper()
    if normalized in {"S", "SÁNG", "SANG"} or "SÁNG" in normalized or "SANG" in normalized:
        return "S"
    if normalized in {"C", "CHIỀU", "CHIEU"} or "CHIỀU" in normalized or "CHIEU" in normalized:
        return "C"
    return value.strip()


def _build_xlsx_class_session_chunks(path: Path) -> list[dict]:
    """
    Chuyển một bảng TKB dạng ma trận thành bản ghi theo:
        1 lớp + 1 ngày + 1 buổi -> 1 chunk
    Mỗi chunk chứa toàn bộ các tiết của lớp trong buổi đó.

    Đây là cấu trúc phù hợp cho truy vấn TKB hơn việc tạo 1 chunk / 1 ô hoặc
    1 chunk / 1 tiết.
    """
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    records: list[dict] = []

    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            title = ""
            for value in rows[0]:
                if value:
                    title = _excel_text(value)
                    break

            header_row_index = None
            headers: list[str] = []
            for idx, row in enumerate(rows[:10], start=1):
                values = [_excel_text(v) for v in row]
                nonempty = [v for v in values if v]
                if len(nonempty) >= 3 and any(v in {"Thứ", "Buổi", "Tiết"} for v in nonempty):
                    header_row_index = idx
                    headers = values
                    break

            if header_row_index is None:
                continue

            current_day = ""
            current_session = ""

            # key -> {period -> class value}
            grouped: dict[tuple[str, str, str, str], dict[str, str]] = {}
            # key day/session -> observed periods in the source sheet
            observed_periods: dict[tuple[str, str], set[str]] = {}
            first_rows: dict[tuple[str, str, str], int] = {}
            last_rows: dict[tuple[str, str, str], int] = {}

            for idx in range(header_row_index, len(rows)):
                row_number = idx + 1
                row = list(rows[idx])
                day_raw = _excel_text(row[0]) if len(row) >= 1 else ""
                session_raw = _excel_text(row[1]) if len(row) >= 2 else ""
                period = _excel_text(row[2]) if len(row) >= 3 else ""

                if day_raw:
                    current_day = _parse_day_value(day_raw)
                if session_raw:
                    current_session = _parse_session_value(session_raw)
                if not current_day or not current_session or not period:
                    continue

                period_key = re.sub(r"[^0-9A-Za-z.-]", "", period)
                if not period_key:
                    continue

                row_has_class_data = False
                for col_idx in range(3, len(row)):
                    if col_idx >= len(headers):
                        continue
                    class_name = _excel_text(headers[col_idx]).upper()
                    if not _is_class_name(class_name):
                        continue

                    value_text = _excel_text(row[col_idx])
                    if value_text:
                        row_has_class_data = True

                    key = (sheet.title, current_day, current_session, class_name)
                    if value_text:
                        grouped.setdefault(key, {})[period_key] = value_text
                        first_rows[key] = min(first_rows.get(key, row_number), row_number)
                        last_rows[key] = max(last_rows.get(key, row_number), row_number)

                if row_has_class_data:
                    observed_periods.setdefault((sheet.title, current_day, current_session), set()).add(period_key)

            for (sheet_name, day, session, class_name), values in grouped.items():
                periods = sorted(observed_periods.get((sheet_name, day, session), set()) or values.keys(), key=_number_key)
                ordered_values = {period: values.get(period, "[TRỐNG]") for period in periods}

                grade_match = re.match(r"(6|7|8|9|10|11|12)A", class_name, re.IGNORECASE)
                grade = grade_match.group(1) if grade_match else ""

                text_lines = [
                    "Thời khóa biểu: " + title if title else "Thời khóa biểu",
                    f"Trang tính: {sheet_name}",
                    f"Khối: {grade}" if grade else "",
                    f"Lớp: {class_name}",
                    f"Thứ: {_day_display(day)}",
                    f"Buổi: {_session_display(session)}",
                ]
                text_lines.extend(
                    f"Tiết {period}: {ordered_values[period]}"
                    for period in periods
                )

                records.append({
                    "sheet": sheet_name,
                    "day": day,
                    "session": session,
                    "class_name": class_name,
                    "grade": grade,
                    "periods": periods,
                    "values": ordered_values,
                    "row_start": first_rows.get((sheet_name, day, session, class_name), 0),
                    "row_end": last_rows.get((sheet_name, day, session, class_name), 0),
                    "text": "\n".join(line for line in text_lines if line),
                })
    finally:
        workbook.close()

    records.sort(
        key=lambda item: (
            item["sheet"],
            int(item["day"]) if str(item["day"]).isdigit() else 99,
            0 if item["session"] == "S" else 1,
            int(re.search(r"\d+", item["class_name"]).group(0)) if re.search(r"\d+", item["class_name"]) else 99,
            item["class_name"],
        )
    )
    return records


def chunk_xlsx(blocks: list[dict]) -> list[dict]:
    """Tạo chunk TKB theo lớp + ngày + buổi, giữ toàn bộ tiết trong cùng chunk."""
    source_path = _xlsx_source_path(blocks)
    if source_path is None:
        return [
            make_chunk([b], i, b["text"])
            for i, b in enumerate(blocks)
            if clean(b.get("text", ""))
        ]

    rows = _build_xlsx_class_session_chunks(source_path)
    if not rows:
        return []

    by_row: dict[tuple[str, int], dict] = {
        (
            b.get("location", {}).get("sheet", ""),
            int(b.get("location", {}).get("row", 0)),
        ): b
        for b in blocks
        if b.get("location", {}).get("sheet") is not None
        and b.get("location", {}).get("row") is not None
    }

    chunks: list[dict] = []
    for index, row in enumerate(rows):
        source_block = by_row.get((row["sheet"], int(row["row_start"])))
        if source_block is None:
            source_block = blocks[0]

        location = dict(source_block.get("location", {}))
        location.update({
            "sheet": row["sheet"],
            "row_start": row["row_start"],
            "row_end": row["row_end"],
            "day": row["day"],
            "session": row["session"],
            "class_name": row["class_name"],
            "grade": row["grade"],
            "periods": ",".join(row["periods"]),
        })

        chunks.append(
            make_chunk(
                [source_block],
                index,
                row["text"],
                extra_location=location,
            )
        )

    return chunks

def build_chunks(blocks: list[dict]) -> tuple[list[dict], dict]:
    by_file: dict[str, list[dict]] = {}
    for block in blocks:
        by_file.setdefault(block["file_id"], []).append(block)

    all_chunks: list[dict] = []
    file_stats: list[dict] = []

    for file_id, file_blocks in by_file.items():
        file_type = file_blocks[0]["file_type"]

        if file_type == "docx":
            chunks = chunk_docx(file_blocks)
        elif file_type == "pdf":
            chunks = chunk_pdf(file_blocks)
        elif file_type in {"xlsx", "xlsm"}:
            chunks = chunk_xlsx(file_blocks)
        else:
            chunks = [
                make_chunk([b], i, b["text"])
                for i, b in enumerate(file_blocks)
            ]

        for i, chunk in enumerate(chunks):
            chunk["chunk_index"] = i
            chunk["id"] = f"{file_id}-c{i:05d}"

        all_chunks.extend(chunks)
        file_stats.append(
            {
                "file_id": file_id,
                "source": file_blocks[0]["source"],
                "file_type": file_type,
                "blocks": len(file_blocks),
                "chunks": len(chunks),
            }
        )

    manifest = {
        "max_chars": MAX_CHARS,
        "max_table_chars": MAX_TABLE_CHARS,
        "total_blocks": len(blocks),
        "total_chunks": len(all_chunks),
        "files": file_stats,
    }
    return all_chunks, manifest


def save_jsonl(records: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for record in records:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")


def save_json(data: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="School AI - Giai đoạn 2: chunking theo cấu trúc tài liệu."
    )
    parser.add_argument("--input", default="data/processed/blocks.jsonl")
    parser.add_argument("--output", default="data/processed/chunks.jsonl")
    parser.add_argument("--manifest", default="data/processed/chunk_manifest.json")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        raise SystemExit(f"Không tìm thấy: {input_path}")

    print("=" * 60)
    print("SCHOOL AI - BƯỚC 2: STRUCTURE-AWARE CHUNKING")
    print("Chưa tạo embedding / Chroma / gọi AI")
    print("=" * 60)

    blocks = load_jsonl(input_path)
    print(f"Blocks đầu vào : {len(blocks)}")

    chunks, manifest = build_chunks(blocks)

    print(f"Chunks đầu ra  : {len(chunks)}")
    print("\nTheo file:")
    for item in manifest["files"]:
        print(
            f"  {item['source']} | {item['file_type']} | "
            f"{item['blocks']} blocks -> {item['chunks']} chunks"
        )

    # Mandatory data-preservation checks.
    corpus = "\n".join(c["text"] for c in chunks)
    print("\nKiểm tra bắt buộc:")
    checks = {
        "'Trương Thị Thanh Châm'": "Trương Thị Thanh Châm",
        "'HIỆU TRƯỞNG'": "HIỆU TRƯỞNG",
        "'Tien Yen’s Got Talent'": "Tien Yen’s Got Talent",
    }
    for label, needle in checks.items():
        print(f"  {label}: {'CÓ' if needle in corpus else 'KHÔNG'}")

    # Show representative XLSX chunks so we can inspect context before embedding.
    xlsx_chunks = [c for c in chunks if c["file_type"] == "xlsx"]
    print("\nKiểm tra mẫu XLSX:")
    for i, chunk in enumerate(xlsx_chunks[:3], start=1):
        print(f"  --- Mẫu {i} ---")
        print(chunk["text"])

    # Check that meaningless single numeric cells are not being embedded.
    numeric_only = []
    for chunk in xlsx_chunks:
        text = chunk["text"]
        meaningful_lines = [
            line
            for line in text.split("\n")
            if not line.startswith(("Trang tính:", "Tiêu đề:", "Thứ:", "Buổi:", "Tiết:"))
        ]
        if meaningful_lines and all(re.fullmatch(r"[^:]+: ?\d+(?:\.0)?", line) for line in meaningful_lines):
            numeric_only.append(chunk)

    print(
        f"  Chunk XLSX chỉ chứa giá trị số không có ngữ cảnh lớp/môn: {len(numeric_only)}"
    )

    save_jsonl(chunks, Path(args.output))
    save_json(manifest, Path(args.manifest))

    print(f"\nĐã lưu chunks: {args.output}")
    print(f"Đã lưu manifest: {args.manifest}")


if __name__ == "__main__":
    main()
